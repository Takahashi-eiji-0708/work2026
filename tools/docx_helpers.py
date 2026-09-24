"""python-docx の共通ヘルパー（日本語フォント設定・見出し・表）。"""
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Cm

JP_FONT = "游ゴシック"
LATIN_FONT = "Yu Gothic"
NAVY = RGBColor(0x00, 0x4D, 0x88)


def _set_run_font(run, size=None, bold=None, color=None):
    run.font.name = LATIN_FONT
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:eastAsia"), JP_FONT)
    rfonts.set(qn("w:ascii"), LATIN_FONT)
    rfonts.set(qn("w:hAnsi"), LATIN_FONT)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def new_document():
    doc = Document()
    sec = doc.sections[0]
    sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
    sec.left_margin = sec.right_margin = Cm(2.0)
    sec.top_margin = sec.bottom_margin = Cm(2.0)
    for name in ["Normal", "Heading 1", "Heading 2", "Heading 3", "Title", "List Bullet", "List Number"]:
        style = doc.styles[name]
        style.font.name = LATIN_FONT
        rpr = style.element.get_or_add_rPr()
        rfonts = rpr.find(qn("w:rFonts"))
        if rfonts is None:
            rfonts = OxmlElement("w:rFonts")
            rpr.insert(0, rfonts)
        rfonts.set(qn("w:eastAsia"), JP_FONT)
        for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
            if rfonts.get(qn(attr)) is not None:
                del rfonts.attrib[qn(attr)]
    doc.styles["Normal"].font.size = Pt(10.5)
    for name, size in (("Title", 20), ("Heading 1", 15), ("Heading 2", 12.5), ("Heading 3", 11)):
        doc.styles[name].font.size = Pt(size)
        doc.styles[name].font.color.rgb = NAVY
    return doc


def para(doc, text="", size=None, bold=False, color=None, align=None, style=None):
    p = doc.add_paragraph(style=style)
    if text:
        add_inline(p, text, size=size, bold=bold, color=color)
    if align == "center":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif align == "right":
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    return p


def add_inline(p, text, size=None, bold=False, color=None):
    """**太字** と `コード` の簡易インライン書式に対応して run を追加する。"""
    import re

    tokens = re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text)
    for tok in tokens:
        if not tok:
            continue
        if tok.startswith("**") and tok.endswith("**"):
            run = p.add_run(tok[2:-2])
            _set_run_font(run, size, True, color)
        elif tok.startswith("`") and tok.endswith("`"):
            run = p.add_run(tok[1:-1])
            _set_run_font(run, size, bold, color)
            run.font.name = "Consolas"
        else:
            run = p.add_run(tok)
            _set_run_font(run, size, bold, color)
    return p


def heading(doc, text, level=1):
    h = doc.add_heading(level=level)
    run = h.add_run(text)
    _set_run_font(run, color=NAVY)
    return h


def bullet(doc, text, numbered=False):
    p = doc.add_paragraph(style="List Number" if numbered else "List Bullet")
    add_inline(p, text)
    return p


def shade(cell, hex_color):
    tcpr = cell._element.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcpr.append(shd)


def table(doc, rows, header=True, widths=None, font_size=9.5):
    t = doc.add_table(rows=len(rows), cols=len(rows[0]))
    t.style = "Table Grid"
    for i, row in enumerate(rows):
        for j, value in enumerate(row):
            cell = t.cell(i, j)
            cell.text = ""
            p = cell.paragraphs[0]
            lines = str(value).split("\n")
            for k, line in enumerate(lines):
                if k:
                    p = cell.add_paragraph()
                add_inline(p, line, size=font_size, bold=(header and i == 0))
            if header and i == 0:
                shade(cell, "DCE9F5")
    if widths:
        for row in t.rows:
            for j, w in enumerate(widths):
                row.cells[j].width = Cm(w)
    doc.add_paragraph()
    return t
