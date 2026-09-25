"""受講者配布用のファイルだけを集めて ZIP にする。

講師用の情報（解答スライド・解答例・ミニ演習集・理解度テストの解答・講師ガイド・参照実装・
Codex 資料の講師用付録・講師による機能確認）は含めない。
    python tools/build_student_package.py          # dist/ に出力
必要なもの：python-docx、python-pptx、LibreOffice（soffice）
"""
import re
import shutil
import subprocess
import sys
import tempfile
import zipfile
from pathlib import Path

from pptx import Presentation

sys.path.insert(0, str(Path(__file__).resolve().parent))
from md_to_docx import convert  # noqa: E402

ROOT = Path(__file__).resolve().parent.parent
TITLE = "ソリューション営業のための生成AIセキュア活用実践"
NAME = f"{TITLE}_受講者配布用"
DECK = ROOT / "slides" / f"{TITLE}.pptx"

README = f"""# {TITLE}　受講者配布資料

2日間（各日 9:30〜17:00）の研修で使う資料です。**演習のデータはすべて架空** です。
実際の顧客情報・社内情報は、研修の中では入力しないでください。

## フォルダ

| フォルダ | 内容 | 使う場面 |
|---|---|---|
| 01_テキスト | 受講者テキスト、演習ガイド、セキュア活用ガイドライン、応用ワークシート、事前準備チェックリスト | 講義・演習・持ち帰り |
| 02_講義スライド | 講義スライド（PDF・ミニ演習の解答ページを除く） | 講義中のメモ・復習 |
| 03_演習素材 | 演習1〜12 で使う Word・Excel・テキスト | 演習 |
| 04_開発トラック（希望者のみ） | Codex 段階開発プロンプト集、AGENTS.md、確認用データ | 第10章の演習14 |

## 研修の前に

`01_テキスト/06_事前準備チェックリスト（受講者用）.docx` の「必須」を、研修の1週間前までに確認してください。
開発トラックを希望する方は「開発トラック希望者のみ」も確認してください。

## 演習素材の使い方

- `03_演習素材/` のファイルは、研修当日に **OneDrive の「研修」フォルダにコピー** してから使います（Copilot を使うため）
- 演習の解答例は、演習の後に講師から配布します

## 研修の後に

- `04_セキュア活用ガイドライン` は職場で使う前のチェックシートです。職場では **自社のガイドラインを優先** してください
- `08_応用ワークシート` は、研修後も更新しながら使ってください
"""


def student_deck_pdf(out_pdf: Path, work: Path):
    """解答スライドを除いた講義スライドを PDF にする。"""
    prs = Presentation(DECK)
    ids = prs.slides._sldIdLst
    removed = 0
    for sld_id, slide in list(zip(list(ids), prs.slides)):
        title = slide.shapes.title.text if slide.shapes.title is not None else ""
        if "解答" in title:
            ids.remove(sld_id)
            removed += 1
    tmp = work / f"{TITLE}（受講者用）.pptx"
    prs.save(tmp)
    subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", str(work), str(tmp)],
                   check=True, capture_output=True, timeout=600)
    shutil.move(str(tmp.with_suffix(".pdf")), out_pdf)
    return len(ids), removed


def env_checklist_md(work: Path) -> Path:
    """環境準備チェックリストから、受講者の事前準備（1章）だけを取り出す。"""
    src = (ROOT / "docs" / "06_環境準備チェックリスト.md").read_text(encoding="utf-8")
    part = src[src.index("## 1. 受講者の事前準備"):src.index("## 2. ")]
    part = part.replace("## 1. 受講者の事前準備（研修の1週間前までに案内）", "## 研修の1週間前までに確認してください")
    md = work / "06_事前準備チェックリスト（受講者用）.md"
    md.write_text("# 事前準備チェックリスト（受講者用）\n\n" + part, encoding="utf-8")
    return md


def codex_md(work: Path) -> Path:
    """Codex 開発プロンプト集から講師用の付録を除き、コピー元のパスを配布フォルダに合わせる。"""
    src = (ROOT / "codex" / "Codex開発プロンプト集.md").read_text(encoding="utf-8")
    src = src[:src.index("## 付録B")].rstrip() + "\n"
    src = src.replace("> - 完成形（参照実装）は `app/` にあります。遅れた場合は、講師の指示で該当するファイルをコピーして追いつきます",
                      "> - 遅れた場合は、講師から完成形のファイルを受け取って追いつきます")
    dist = f"../{NAME}/04_開発トラック（希望者のみ）"
    src = src.replace("# 研修リポジトリから AGENTS.md と確認用データをコピー", "# 配布フォルダから AGENTS.md と確認用データをコピー")
    src = src.replace("cp ../work2026/app/AGENTS.md .", f"cp {dist}/AGENTS.md .")
    src = src.replace("cp ../work2026/codex/testdata/*.txt testdata/", f"cp {dist}/testdata/*.txt testdata/")
    src = src.replace("cp ../work2026/materials/03_商談文字起こし_初回訪問.txt testdata/",
                      f"cp ../{NAME}/03_演習素材/03_商談文字起こし_初回訪問.txt testdata/")
    leftover = re.findall(r"講師用|app/tests|work2026", src)
    if leftover:
        raise SystemExit(f"受講者用の Codex 資料に講師用の記述が残っています: {leftover}")
    md = work / "Codex開発プロンプト集.md"
    md.write_text(src, encoding="utf-8")
    return md


def main():
    out_dir = ROOT / "dist"
    out_dir.mkdir(exist_ok=True)
    with tempfile.TemporaryDirectory() as t:
        work = Path(t)
        pkg = work / NAME
        text, slides, mats, dev = (pkg / "01_テキスト", pkg / "02_講義スライド", pkg / "03_演習素材",
                                   pkg / "04_開発トラック（希望者のみ）")
        for d in (text, slides, mats, dev / "testdata"):
            d.mkdir(parents=True)

        readme = work / "00_はじめにお読みください.md"
        readme.write_text(README, encoding="utf-8")
        convert(readme, pkg / "00_はじめにお読みください.docx")

        for name in ("02_受講者テキスト", "03_演習ガイド", "04_セキュア活用ガイドライン", "08_応用ワークシート"):
            shutil.copy(ROOT / "docs" / "word" / f"{name}.docx", text)
        md = env_checklist_md(work)
        convert(md, text / (md.stem + ".docx"))

        n, removed = student_deck_pdf(slides / f"{TITLE}（受講者用）.pdf", work)

        for f in sorted((ROOT / "materials").iterdir()):
            if f.is_file():
                shutil.copy(f, mats)

        md = codex_md(work)
        convert(md, dev / "Codex開発プロンプト集.docx")
        shutil.copy(md, dev)
        shutil.copy(ROOT / "app" / "AGENTS.md", dev)
        for f in sorted((ROOT / "codex" / "testdata").glob("*.txt")):
            shutil.copy(f, dev / "testdata")

        out = out_dir / f"{NAME}.zip"
        with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
            for f in sorted(pkg.rglob("*")):
                if f.is_file():
                    z.write(f, f.relative_to(work))
        with zipfile.ZipFile(out) as z:
            names = z.namelist()
            assert z.testzip() is None
            bad = [x for x in names if "解答" in x or "講師" in x or "ミニ演習集" in x or "理解度" in x]
            if bad:
                raise SystemExit(f"講師用のファイルが含まれています: {bad}")
        print(f"{out.relative_to(ROOT)}（{len(names)}ファイル、スライド {n}枚・解答スライド {removed}枚を除外）")
        for x in names:
            print("  " + x)


if __name__ == "__main__":
    main()
